"""Рендер резюме в DOCX и PDF.

Формат нарочно скучный: одна колонка, стандартные заголовки, никаких таблиц,
врезок, рамок и иконок — всё это ATS читает плохо или не читает вовсе.
Шрифт обычный системный, кегль читаемый, скрытого текста нет и быть не может:
рендер умеет писать только строки `Resume`, а они собираются из фактов.
"""

from __future__ import annotations

import html
from pathlib import Path

from .resume import Resume

# Шрифты с кириллицей, которые обычно есть под рукой. Первый найденный и берём.
FONT_CANDIDATES: tuple[tuple[str, str], ...] = (
    ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
    ("C:/Windows/Fonts/calibri.ttf", "C:/Windows/Fonts/calibrib.ttf"),
    (
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
    ),
    (
        "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
        "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf",
    ),
)

DOCX_FONT = "Arial"


class RenderError(RuntimeError):
    """Нечем или некуда рендерить."""


def find_font() -> tuple[Path, Path]:
    for regular, bold in FONT_CANDIDATES:
        regular_path, bold_path = Path(regular), Path(bold)
        if regular_path.exists() and bold_path.exists():
            return regular_path, bold_path
    raise RenderError(
        "не нашёлся шрифт с кириллицей для PDF. Проверенные пути: "
        + ", ".join(r for r, _ in FONT_CANDIDATES)
    )


def render_docx(resume: Resume, path: Path) -> Path:
    from docx import Document
    from docx.shared import Pt

    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = DOCX_FONT
    normal.font.size = Pt(11)

    name = document.add_paragraph()
    run = name.add_run(resume.name.text)
    run.bold = True
    run.font.size = Pt(18)

    for section in resume.sections:
        document.add_paragraph(section.heading, style="Heading 1")
        for line in section.lines:
            if line.kind == "item":
                paragraph = document.add_paragraph(style="List Bullet")
                paragraph.add_run(line.text).bold = True
            elif line.kind == "subtitle":
                paragraph = document.add_paragraph()
                paragraph.add_run(line.text).bold = True
            else:
                document.add_paragraph(line.text)

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    return path


def render_pdf(resume: Resume, path: Path) -> Path:
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    regular, bold = find_font()
    pdfmetrics.registerFont(TTFont("ResumeFont", str(regular)))
    pdfmetrics.registerFont(TTFont("ResumeFont-Bold", str(bold)))

    base = ParagraphStyle(
        "body",
        fontName="ResumeFont",
        fontSize=10.5,
        leading=14,
        alignment=TA_LEFT,
        spaceAfter=2,
    )
    styles = {
        "name": ParagraphStyle(
            "name", parent=base, fontName="ResumeFont-Bold", fontSize=18, leading=22,
            spaceAfter=6,
        ),
        "heading": ParagraphStyle(
            "heading", parent=base, fontName="ResumeFont-Bold", fontSize=12.5,
            leading=16, spaceBefore=10, spaceAfter=4,
        ),
        "subtitle": ParagraphStyle(
            "subtitle", parent=base, fontName="ResumeFont-Bold", fontSize=11
        ),
        "item": ParagraphStyle(
            "item", parent=base, fontName="ResumeFont-Bold", spaceBefore=5
        ),
        "detail": ParagraphStyle("detail", parent=base, leftIndent=10),
    }

    path.parent.mkdir(parents=True, exist_ok=True)
    document = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title=resume.name.text,
        author=resume.name.text,
    )

    flowables = [Paragraph(html.escape(resume.name.text), styles["name"])]
    for section in resume.sections:
        flowables.append(Paragraph(html.escape(section.heading), styles["heading"]))
        for line in section.lines:
            style = styles.get(line.kind, base)
            flowables.append(Paragraph(html.escape(line.text), style))
    flowables.append(Spacer(1, 1))

    document.build(flowables)
    return path
