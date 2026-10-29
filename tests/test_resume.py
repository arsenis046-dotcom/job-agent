"""Главная проверка Фазы 1: в резюме не может оказаться ничего,
чего нет в мастер-профиле."""

from __future__ import annotations

import pytest

from jobagent.documents import (
    HEADINGS,
    LABELS,
    Line,
    ResumeError,
    build_resume,
    check_pdf,
    extract_text,
    render_docx,
    render_pdf,
)
from jobagent.documents.ats import normalize
from jobagent.profile import load_master_profile


def profile_facts(profile) -> set[str]:
    """Все строковые значения профиля — множество допустимых фактов."""
    facts: set[str] = set()

    def walk(value) -> None:
        if isinstance(value, dict):
            for item in value.values():
                walk(item)
        elif isinstance(value, list):
            for item in value:
                walk(item)
        elif value is not None and not isinstance(value, bool):
            facts.add(str(value).strip())

    walk(profile.model_dump(mode="json"))
    return facts


def test_каждый_факт_в_резюме_взят_из_профиля(sample_profile):
    document = build_resume(sample_profile)
    allowed = profile_facts(sample_profile)
    for line in document.all_lines:
        for fact in line.evidence:
            assert fact in allowed, f"{line.source}: факта «{fact}» нет в профиле"


def test_в_резюме_нет_текста_кроме_фактов_и_ярлыков(sample_profile):
    """Текст строки — вычисляемое свойство: ярлык из закрытого списка плюс факты."""
    document = build_resume(sample_profile)
    for line in document.all_lines:
        residue = line.text
        if line.label:
            assert line.label in LABELS
            residue = residue.replace(f"{line.label}: ", "", 1)
        for fact in line.evidence:
            residue = residue.replace(fact, "", 1)
        assert not residue.strip(" ,—-"), f"{line.source}: лишний текст «{residue}»"


def test_выдуманный_ярлык_не_собирается():
    with pytest.raises(ResumeError):
        Line(kind="detail", evidence=("Python",), source="x", label="Опыт работы")


def test_строка_без_фактов_не_собирается():
    with pytest.raises(ResumeError):
        Line(kind="text", evidence=("",), source="x")


def test_без_обязательных_фактов_резюме_не_собирается(repo):
    """Пустой профиль из репозитория: вместо пустого резюме — список вопросов."""
    master = load_master_profile(repo)
    with pytest.raises(ResumeError) as exc:
        build_resume(master)
    assert "person.full_name" in str(exc.value)


def test_разделы_стандартные_и_в_нужном_порядке(sample_profile):
    document = build_resume(sample_profile)
    assert document.headings == [h for h in HEADINGS if h in document.headings]
    for heading in document.headings:
        assert heading in HEADINGS, "самодельных заголовков ATS не ждёт"
    assert "Проекты" in document.headings
    assert "Опыт работы" not in document.headings, (
        "у этого профиля опыта работы нет — пустого раздела быть не должно"
    )


def test_опыт_работы_попадает_в_резюме(sample_profile):
    from jobagent.profile import Job

    профиль = sample_profile.model_copy(
        update={
            "employment_history": [
                Job(
                    position="Специалист технической поддержки",
                    company="ООО Пример",
                    period="январь 2025 — настоящее время",
                    duties=["приём и решение заявок"],
                    achievements=["внедрил внутренний сервис"],
                )
            ]
        }
    )
    document = build_resume(профиль)
    assert document.headings.index("Опыт работы") < document.headings.index("Проекты")

    раздел = next(s for s in document.sections if s.heading == "Опыт работы")
    тексты = [line.text for line in раздел.lines]
    assert тексты[0] == "Специалист технической поддержки — ООО Пример"
    assert "Период: январь 2025 — настоящее время" in тексты
    assert "Обязанности: приём и решение заявок" in тексты
    assert "Сделано: внедрил внутренний сервис" in тексты


def test_где_сделан_проект_видно_сразу_после_названия(sample_profile):
    """Работодатель спрашивает об этом первым."""
    project = sample_profile.projects[0].model_copy(
        update={"context": "рабочая задача в ООО Пример"}
    )
    document = build_resume(sample_profile.model_copy(update={"projects": [project]}))
    проекты = next(s for s in document.sections if s.heading == "Проекты")
    assert проекты.lines[1].label == "Где сделано"
    assert проекты.lines[1].text == "Где сделано: рабочая задача в ООО Пример"


def test_проект_подан_как_проект(sample_profile):
    document = build_resume(sample_profile)
    projects = next(s for s in document.sections if s.heading == "Проекты")
    labels = [line.label for line in projects.lines if line.label]
    assert labels == ["Стек", "Моя роль", "Результат", "Ссылка"]


def test_docx_собирается_одной_колонкой(sample_profile, tmp_path):
    from docx import Document

    path = render_docx(build_resume(sample_profile), tmp_path / "r.docx")
    document = Document(path)
    assert not document.tables, "таблиц в резюме быть не должно"
    text = "\n".join(p.text for p in document.paragraphs)
    assert "Иван Петров" in text
    assert "Стек: Python, aiogram, SQLite" in text


def test_pdf_читается_как_текст_и_порядок_не_поехал(sample_profile, tmp_path):
    document = build_resume(sample_profile)
    path = render_pdf(document, tmp_path / "r.pdf")
    report = check_pdf(path, document)
    assert report.ok, str(report)
    assert report.checked_lines == len(document.all_lines)


def test_проверка_atс_замечает_подмену(sample_profile, tmp_path):
    """Если в PDF попало не то, что в документе, проверка обязана это увидеть."""
    document = build_resume(sample_profile)
    path = render_pdf(document, tmp_path / "r.pdf")

    подменённый = build_resume(
        sample_profile.model_copy(
            update={
                "person": sample_profile.person.model_copy(
                    update={"full_name": "Пётр Иванов"}
                )
            }
        )
    )
    report = check_pdf(path, подменённый)
    assert not report.ok
    assert any("имя не читается" in p for p in report.problems)


def test_в_pdf_нет_скрытого_текста(sample_profile, tmp_path):
    """Извлечённый текст не должен содержать ничего сверх строк документа."""
    document = build_resume(sample_profile)
    path = render_pdf(document, tmp_path / "r.pdf")
    extracted = normalize(extract_text(path))

    known = normalize(document.text() + " " + " ".join(document.headings))
    for word in extracted.split():
        assert word in known, f"в PDF нашлось лишнее слово: {word}"
